import csv
import io
import re
import zipfile
from pathlib import Path
from urllib.parse import urlparse
from uuid import uuid4

import httpx
import openpyxl
import pypdf
from docx import Document as DocxDocument
from fastapi import UploadFile
from pptx import Presentation
from youtube_transcript_api import YouTubeTranscriptApi

from app.config import settings
from app.models.product import ReferenceAnalysis, ReferenceKind, ReferenceMaterial
from app.services.openrouter_utils import call_openrouter_with_retry, generate_text, image_part

_CHUNK_SIZE = 1024 * 1024
_MAX_DOC_BYTES = settings.max_reference_upload_mb * 1024 * 1024
_MAX_MEDIA_BYTES = _MAX_DOC_BYTES * 4
_MAX_MATERIAL_TEXT_CHARS = 40_000
_MAX_URL_TEXT_CHARS = 200_000
_MAX_ZIP_ENTRIES = 50
_MAX_ZIP_UNCOMPRESSED_BYTES = 60 * 1024 * 1024

_JINA_BASE = "https://r.jina.ai/"
_YOUTUBE_ID_RE = re.compile(r"(?:v=|youtu\.be/|embed/)([A-Za-z0-9_-]{11})")

_EXTENSION_KIND = {
    ".pdf": ReferenceKind.pdf,
    ".docx": ReferenceKind.docx,
    ".doc": ReferenceKind.doc,
    ".pptx": ReferenceKind.pptx,
    ".ppt": ReferenceKind.ppt,
    ".txt": ReferenceKind.txt,
    ".csv": ReferenceKind.csv,
    ".xlsx": ReferenceKind.xlsx,
    ".zip": ReferenceKind.zip,
    ".jpg": ReferenceKind.image,
    ".jpeg": ReferenceKind.image,
    ".png": ReferenceKind.image,
    ".gif": ReferenceKind.image,
    ".webp": ReferenceKind.image,
    ".mp4": ReferenceKind.video,
    ".mov": ReferenceKind.video,
    ".webm": ReferenceKind.video,
    ".avi": ReferenceKind.video,
    ".mkv": ReferenceKind.video,
    ".mp3": ReferenceKind.audio,
    ".wav": ReferenceKind.audio,
    ".m4a": ReferenceKind.audio,
    ".aac": ReferenceKind.audio,
    ".ogg": ReferenceKind.audio,
    ".flac": ReferenceKind.audio,
}

_IMAGE_MEDIA_TYPES = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".gif": "image/gif",
    ".webp": "image/webp",
}

_IMAGE_DESCRIBE_SYSTEM_PROMPT = """Describe this image factually for a product-marketing team building
an ad from it. Note the product, branding/packaging, setting, colors, and any visible text — 2-4 sentences.
Don't speculate about facts you can't see."""

_NOT_SUPPORTED_NOTE = (
    "Legacy .doc/.ppt files aren't supported for analysis yet — convert to .docx/.pptx for the AI to "
    "read it. The file stays attached for reference only."
)
_COMING_SOON_NOTE = "Video/audio analysis is coming soon — stored for preview only."


class ExtractionError(Exception):
    """Raised for any user-facing extraction failure (bad URL, unsupported type, oversized file)."""


def _kind_from_extension(ext: str) -> ReferenceKind | None:
    return _EXTENSION_KIND.get(ext.lower())


def _cap_text(text: str) -> tuple[str, bool]:
    text = text.strip()
    if len(text) > _MAX_MATERIAL_TEXT_CHARS:
        return text[:_MAX_MATERIAL_TEXT_CHARS], True
    return text, False


def _describe_image_with_gemini(data: bytes, media_type: str) -> str:
    """Never raises — a failed vision call just yields an empty description
    rather than blocking the whole upload."""
    try:
        text = call_openrouter_with_retry(
            lambda: generate_text(
                system_instruction=_IMAGE_DESCRIBE_SYSTEM_PROMPT,
                contents=[image_part(data, media_type), "Describe this image."],
                model=settings.openrouter_text_model,
                max_output_tokens=512,
            ),
            label="describe_image",
        )
        return text.strip()
    except Exception:
        return ""


def _extract_pdf(data: bytes) -> str:
    reader = pypdf.PdfReader(io.BytesIO(data))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages[:200])


def _extract_docx(data: bytes) -> str:
    doc = DocxDocument(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_pptx(data: bytes) -> str:
    prs = Presentation(io.BytesIO(data))
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
                if text.strip():
                    parts.append(text)
    return "\n\n".join(parts)


def _extract_csv(data: bytes) -> str:
    rows = list(csv.reader(io.StringIO(data.decode("utf-8", errors="replace"))))[:500]
    return "\n".join(", ".join(row) for row in rows)


def _extract_xlsx(data: bytes) -> str:
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"[Sheet: {sheet.title}]")
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i >= 500:
                break
            parts.append(", ".join("" if v is None else str(v) for v in row))
    return "\n".join(parts)


def _extract_bytes_by_kind(kind: ReferenceKind, ext: str, data: bytes) -> str:
    if kind == ReferenceKind.pdf:
        return _extract_pdf(data)
    if kind == ReferenceKind.docx:
        return _extract_docx(data)
    if kind == ReferenceKind.pptx:
        return _extract_pptx(data)
    if kind == ReferenceKind.txt:
        return data.decode("utf-8", errors="replace")
    if kind == ReferenceKind.csv:
        return _extract_csv(data)
    if kind == ReferenceKind.xlsx:
        return _extract_xlsx(data)
    if kind == ReferenceKind.image:
        return _describe_image_with_gemini(data, _IMAGE_MEDIA_TYPES.get(ext, "image/jpeg"))
    return ""


def _extract_zip(data: bytes) -> tuple[str, str]:
    blocks: list[str] = []
    skipped: list[str] = []
    total_uncompressed = 0
    _UNSUPPORTED_IN_ZIP = {ReferenceKind.video, ReferenceKind.audio, ReferenceKind.zip, ReferenceKind.doc, ReferenceKind.ppt}

    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        infos = [i for i in zf.infolist() if not i.is_dir()]
        for info in infos[:_MAX_ZIP_ENTRIES]:
            total_uncompressed += info.file_size
            if total_uncompressed > _MAX_ZIP_UNCOMPRESSED_BYTES:
                skipped.append(f"{info.filename} (zip size cap reached)")
                break

            ext = Path(info.filename).suffix.lower()
            kind = _kind_from_extension(ext)
            if kind is None or kind in _UNSUPPORTED_IN_ZIP:
                skipped.append(info.filename)
                continue

            try:
                text = _extract_bytes_by_kind(kind, ext, zf.read(info))
            except Exception:
                skipped.append(f"{info.filename} (failed to read)")
                continue

            if text.strip():
                blocks.append(f"--- {info.filename} ---\n{text}")

        if len(infos) > _MAX_ZIP_ENTRIES:
            skipped.append(f"...and {len(infos) - _MAX_ZIP_ENTRIES} more entries (entry cap reached)")

    note = f"Skipped {len(skipped)} unsupported/oversized entries in this zip." if skipped else ""
    return "\n\n".join(blocks), note


async def extract_from_upload(file: UploadFile) -> ReferenceMaterial:
    """Stream an uploaded file to disk, then dispatch extraction by extension.
    Video/audio are stored + previewable only (analysis explicitly out of scope
    for this pass); legacy .doc/.ppt are stored only (no pure-Python parser)."""

    filename = file.filename or "upload"
    ext = Path(filename).suffix.lower()
    kind = _kind_from_extension(ext)
    if kind is None:
        raise ExtractionError(f"Unsupported file type '{ext or filename}'.")

    is_media = kind in (ReferenceKind.video, ReferenceKind.audio)
    size_cap = _MAX_MEDIA_BYTES if is_media else _MAX_DOC_BYTES

    material_id = uuid4().hex[:12]
    uploads_dir = Path(settings.reference_uploads_dir).resolve()
    uploads_dir.mkdir(parents=True, exist_ok=True)
    stored_name = f"{material_id}{ext}"
    stored_path = uploads_dir / stored_name

    total = 0
    buffer = bytearray()
    with open(stored_path, "wb") as out:
        while chunk := await file.read(_CHUNK_SIZE):
            total += len(chunk)
            if total > size_cap:
                out.close()
                stored_path.unlink(missing_ok=True)
                raise ExtractionError(
                    f"'{filename}' is larger than the {size_cap // (1024 * 1024)}MB limit for this file type."
                )
            out.write(chunk)
            if not is_media:
                buffer.extend(chunk)

    material = ReferenceMaterial(
        id=material_id,
        kind=kind,
        filename=filename,
        stored_path=f"reference_uploads/{stored_name}",
        mime_type=file.content_type,
        size_bytes=total,
        analysis=ReferenceAnalysis.failed,
    )

    if is_media:
        material.analysis = ReferenceAnalysis.coming_soon
        material.note = _COMING_SOON_NOTE
        return material

    if kind in (ReferenceKind.doc, ReferenceKind.ppt):
        material.analysis = ReferenceAnalysis.not_supported
        material.note = _NOT_SUPPORTED_NOTE
        return material

    try:
        if kind == ReferenceKind.zip:
            text, zip_note = _extract_zip(bytes(buffer))
            material.note = zip_note
        else:
            text = _extract_bytes_by_kind(kind, ext, bytes(buffer))
    except Exception as e:
        material.analysis = ReferenceAnalysis.failed
        material.note = f"Couldn't extract content from this file: {e}"
        return material

    capped_text, truncated = _cap_text(text)
    material.extracted_text = capped_text
    material.truncated = truncated
    if capped_text.strip():
        material.analysis = ReferenceAnalysis.analyzed
    else:
        material.analysis = ReferenceAnalysis.failed
        material.note = material.note or "No readable text found in this file."
    return material


def classify_link_kind(url: str) -> ReferenceKind:
    """Hostname sniff for card icon/label only — extraction itself only
    branches YouTube vs. everything else (the generic Jina Reader path)."""
    host = urlparse(url).netloc.lower()
    if "youtube.com" in host or "youtu.be" in host:
        return ReferenceKind.youtube
    if "drive.google.com" in host or "docs.google.com" in host:
        return ReferenceKind.google_drive
    if "dropbox.com" in host:
        return ReferenceKind.dropbox
    if "notion.so" in host or "notion.site" in host:
        return ReferenceKind.notion
    return ReferenceKind.website


def _fetch_via_jina(url: str) -> tuple[str, str]:
    try:
        resp = httpx.get(
            f"{_JINA_BASE}{url}",
            headers={"X-Return-Format": "markdown"},
            timeout=settings.jina_reader_timeout_seconds,
            follow_redirects=True,
        )
        resp.raise_for_status()
    except httpx.HTTPStatusError as e:
        raise ExtractionError(
            f"Couldn't fetch that page (HTTP {e.response.status_code}). Double-check the URL is correct and public."
        ) from e
    except httpx.RequestError as e:
        raise ExtractionError("Couldn't reach that URL — check it's correct and publicly accessible.") from e

    text = resp.text[:_MAX_URL_TEXT_CHARS]
    if not text.strip():
        raise ExtractionError("That page returned no readable content.")

    title = ""
    first_line = text.strip().splitlines()[0]
    if first_line.lower().startswith("title:"):
        title = first_line.split(":", 1)[1].strip()
    return text, title


def _fetch_youtube_transcript(url: str) -> tuple[str, str]:
    match = _YOUTUBE_ID_RE.search(url)
    if not match:
        raise ExtractionError("Couldn't find a video ID in that YouTube URL.")
    try:
        segments = YouTubeTranscriptApi.get_transcript(match.group(1))
    except Exception as e:
        raise ExtractionError(
            "Couldn't get a transcript for that video — captions may be disabled, or it's unavailable."
        ) from e
    text = " ".join(seg["text"] for seg in segments)
    return text[:_MAX_URL_TEXT_CHARS], ""


def fetch_url_as_text(url: str) -> tuple[str, str]:
    """Returns (raw_text, title). Dispatches YouTube transcripts vs. the
    generic Jina Reader path used for websites and any Drive/Dropbox/Notion
    link pasted as a plain URL (no OAuth integration in this pass)."""
    if classify_link_kind(url) == ReferenceKind.youtube:
        return _fetch_youtube_transcript(url)
    return _fetch_via_jina(url)
